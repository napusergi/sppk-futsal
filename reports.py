"""Export laporan perhitungan AHP & TOPSIS ke PDF (reportlab) dan Excel (openpyxl)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import sppk
from sppk import POSITIONS

COURT = colors.HexColor("#0f6b57")

# Font TTF agar simbol λ, Σ, √, ≥ tampil; fallback Helvetica dengan simbol ASCII
FONT, FONT_B = "Helvetica", "Helvetica-Bold"
try:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
    pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))
    FONT, FONT_B = "Arial", "Arial-Bold"
except Exception:
    pass


def clean(v):
    if not isinstance(v, str):
        return v
    v = v.replace("⁺", "+").replace("⁻", "-").replace("ₚ", "p").replace("ₖ", "k")
    if FONT == "Helvetica":
        for a, b in (("λ", "lambda "), ("Σ", "sum "), ("√", "sqrt "), ("≥", ">="), ("→", "->"), ("—", "-"), ("−", "-")):
            v = v.replace(a, b)
    return v


def P(text, style):
    return Paragraph(clean(text), style)
SOFT = colors.HexColor("#e1efea")


def sheet_name(text, existing=()):
    """Nama sheet Excel: tanpa karakter terlarang, maks 31 karakter, unik."""
    base = "".join("-" if ch in '\\/:*?[]' else ch for ch in text)[:31]
    name, k = base, 2
    while name in existing:
        name = f"{base[:28]}~{k}"
        k += 1
    return name


def frac(v):
    return str(int(round(v))) if v >= 1 else f"1/{int(round(1 / v))}"


def codes(criteria):
    return [f"C{i + 1}" for i in range(len(criteria))]


# ======================================================================= PDF
def _styles():
    ss = getSampleStyleSheet()
    return {"h1": ParagraphStyle("h1", parent=ss["Title"], fontName=FONT_B, fontSize=16, textColor=COURT, spaceAfter=4),
            "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=FONT_B, fontSize=12, textColor=COURT, spaceBefore=8, spaceAfter=4),
            "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontName=FONT_B, fontSize=9.5, spaceBefore=6, spaceAfter=3),
            "p": ParagraphStyle("p", parent=ss["BodyText"], fontName=FONT, fontSize=8.5, leading=11),
            "small": ParagraphStyle("s", parent=ss["BodyText"], fontName=FONT, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#5a6b66"))}


def _table(rows, header_rows=1, first_col=True, font=6.8, col_widths=None, hl_rows=()):
    rows = [[clean(v) for v in r] for r in rows]
    t = Table(rows, colWidths=col_widths, repeatRows=header_rows)
    st = [("FONT", (0, 0), (-1, -1), FONT, font),
          ("FONT", (0, 0), (-1, header_rows - 1), FONT_B, font),
          ("BACKGROUND", (0, 0), (-1, header_rows - 1), SOFT),
          ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b8c7c2")),
          ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
    if first_col:
        st.append(("FONT", (0, header_rows), (0, -1), FONT_B, font))
    for r in hl_rows:
        st += [("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f6ead9")), ("FONT", (0, r), (-1, r), FONT_B, font)]
    t.setStyle(TableStyle(st))
    return t


def _example_table(steps, S):
    cell = ParagraphStyle("ex", parent=S["p"], fontSize=7.5, leading=9.5)
    rows = [[P("<b>Langkah</b>", cell), P("<b>Rumus</b>", cell), P("<b>Substitusi angka</b>", cell)]]
    rows += [[P(f"<b>{t}</b>", cell), P(f, cell), P(x, cell)] for t, f, x in steps]
    t = Table(rows, colWidths=[4.2 * cm, 6.8 * cm, 16.2 * cm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b8c7c2")), ("BACKGROUND", (0, 0), (-1, 0), SOFT),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _doc(title):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.1 * cm, bottomMargin=1.1 * cm, title=title)

    def footer(c, d):
        c.saveState()
        c.setFont(FONT, 7)
        c.setFillColor(colors.grey)
        c.drawString(1.2 * cm, 0.6 * cm, f"SPK Hibrida AHP–TOPSIS Futsal · {title} · dicetak {datetime.now():%d-%m-%Y %H:%M}")
        c.drawRightString(landscape(A4)[0] - 1.2 * cm, 0.6 * cm, f"Halaman {d.page}")
        c.restoreState()
    return buf, doc, footer


def _legend(criteria, S):
    rows = [["Kode", "Kriteria", "Kode", "Kriteria", "Kode", "Kriteria"]]
    cells = [(f"C{i + 1}", c["name"]) for i, c in enumerate(criteria)]
    for i in range(0, len(cells), 3):
        chunk = cells[i:i + 3] + [("", "")] * (3 - len(cells[i:i + 3]))
        rows.append([x for pair in chunk for x in pair])
    t = _table(rows, font=7.5, first_col=False)
    t.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "LEFT")]))
    return t


def ahp_pdf(criteria, profiles):
    S = _styles()
    buf, doc, footer = _doc("Perhitungan AHP")
    cc = codes(criteria)
    el = [P("Laporan Perhitungan AHP — Bobot Kriteria per Posisi", S["h1"]),
          P("Hierarki: <b>Tujuan</b> menentukan posisi terbaik pemain futsal → <b>Kriteria</b> "
                    f"{len(criteria)} skill → <b>Alternatif</b> {', '.join(POSITIONS)}.", S["p"]),
          P("Model: hibrida AHP–TOPSIS dengan analisis gap berbobot; AHP menghasilkan bobot kriteria per posisi. "
            "Matriks kepentingan antar-kriteria ditetapkan <b>sekali oleh pakar</b> (pelatih) per posisi pada skala Saaty 1–9; "
            "setiap pemain baru dinilai memakai matriks tersimpan ini. Matriks berlabel <b>contoh</b> adalah nilai awal untuk simulasi yang "
            "harus diganti penilaian pakar. Nilai contoh dibentuk dari "
            "skor kebutuhan 1–9 (High = 9, Medium = 5, Low = 1): d = s<sub>i</sub> − s<sub>j</sub>; "
                    "a<sub>ij</sub> = d + 1 (d ≥ 0), a<sub>ij</sub> = 1/(1 − d) (d &lt; 0), dibatasi 1–9. "
                    "Bobot w<sub>i</sub> = rata-rata baris matriks ternormalisasi (a<sub>ij</sub>/Σ<sub>i</sub>a<sub>ij</sub>). "
                    "λmax = rata-rata (Aw)<sub>i</sub>/w<sub>i</sub>; CI = (λmax − n)/(n − 1); CR = CI/RI; konsisten bila CR &lt; 0,1.", S["p"]),
          Spacer(1, 4), _legend(criteria, S)]
    # ringkasan
    summ = [["Kriteria"] + POSITIONS]
    al = [sppk.ahp_pos(profiles, p) for p in range(len(POSITIONS))]
    for i, c in enumerate(criteria):
        summ.append([f"{cc[i]} {c['name']}"] + [f"{a['weights'][i]:.4f}" for a in al])
    summ += [["λmax"] + [f"{a['lambda_max']:.4f}" for a in al], ["CI"] + [f"{a['ci']:.4f}" for a in al],
             ["CR"] + [f"{a['cr']:.4f}" for a in al], ["Status"] + ["Konsisten" if a["consistent"] else "Tidak konsisten" for a in al],
             ["Sumber matriks"] + [{"pakar": "Penilaian pakar", "contoh": "Contoh (belum diisi pakar)"}.get(getattr(profiles, "sources", {}).get(k), "Belum diisi (semua = 1)") for k in range(len(al))]]
    el += [P("Ringkasan bobot semua posisi", S["h2"]), _table(summ, font=7.5, hl_rows=(len(summ) - 2,))]

    for p, pos in enumerate(POSITIONS):
        a = al[p]
        s = profiles[p]
        n = a["n"]
        el += [PageBreak(), P(f"AHP — {pos}", S["h2"]),
               P("Hasil: bobot prioritas kriteria (vektor prioritas AHP)", S["h3"]),
               _table([["Kode", "Kriteria", "Bobot", "%", "Peringkat"]] +
                      [[cc[i], c["name"], f"{a['weights'][i]:.4f}", f"{a['weights'][i] * 100:.1f}%",
                        str(1 + sum(1 for x in a["weights"] if x > a["weights"][i]))] for i, c in enumerate(criteria)],
                      font=7, first_col=False, col_widths=[1.6 * cm, 5 * cm, 2.4 * cm, 2 * cm, 2.2 * cm])]
        el += [P("Langkah 1. Matriks perbandingan berpasangan — sumber: " + {"pakar": "Penilaian pakar", "contoh": "Contoh (belum diisi pakar)"}.get(getattr(profiles, "sources", {}).get(p), "Belum diisi (semua = 1)"), S["h3"]),
               _table([[""] + cc] + [[cc[i]] + [frac(v) for v in row] for i, row in enumerate(a["matrix"])] +
                      [["Σ kolom"] + [f"{x:.3f}" for x in a["col_sum"]]], hl_rows=(n + 1,))]
        el += [P("Langkah 2. Normalisasi matriks dan bobot prioritas", S["h3"]),
               _table([[""] + cc + ["Bobot w"]] + [[cc[i]] + [f"{v:.3f}" for v in row] + [f"{a['weights'][i]:.4f}"]
                                                    for i, row in enumerate(a["normalized"])])]
        el += [P("Langkah 3. Uji konsistensi", S["h3"]),
               _table([["Kode", "w", "(Aw)", "(Aw)/w"]] + [[cc[i], f"{a['weights'][i]:.4f}", f"{a['aw'][i]:.4f}", f"{a['ratio'][i]:.4f}"] for i in range(n)],
                      font=7, col_widths=[2 * cm, 3 * cm, 3 * cm, 3 * cm]),
               Spacer(1, 4),
               P(f"λmax = {a['lambda_max']:.4f};  CI = ({a['lambda_max']:.4f} − {n})/({n} − 1) = {a['ci']:.4f};  "
                         f"RI(n={n}) = {a['ri']};  CR = {a['ci']:.4f}/{a['ri']} = <b>{a['cr']:.4f}</b> → "
                         f"<b>{'KONSISTEN' if a['consistent'] else 'TIDAK KONSISTEN'}</b>", S["p"])]
        el += [P(f"Contoh perhitungan: {criteria[0]['name']} vs {criteria[1]['name'] if n > 1 else criteria[0]['name']}", S["h3"]),
               _example_table(sppk.ahp_example([c["name"] for c in criteria], s, a, 0, 1 if n > 1 else 0), S)]
    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


def topsis_pdf(criteria, profiles, players):
    S = _styles()
    buf, doc, footer = _doc("Perhitungan TOPSIS")
    cc = codes(criteria)
    pc = [p["code"] for p in players]
    names = [c["name"] for c in criteria]
    el = [P("Laporan Perhitungan TOPSIS — Kecocokan Posisi Pemain", S["h1"]),
          P("Model hibrida AHP–TOPSIS. <b>Penentu rekomendasi:</b> satu TOPSIS per pemain dengan <b>alternatif = semua posisi</b>, "
            "matriks keputusan x<sub>pj</sub> = bobot AHP kriteria j pada posisi p, bobot kriteria w<sub>j</sub> = s<sub>j</sub>/Σs "
            "(profil skill pemain), semua kriteria benefit. Langkah: r<sub>pj</sub> = x<sub>pj</sub>/√Σx<sub>pj</sub>²; "
            "v<sub>pj</sub> = w<sub>j</sub>·r<sub>pj</sub>; A⁺ = max v, A⁻ = min v; D⁺ = √Σ(v − A⁺)², D⁻ = √Σ(v − A⁻)²; "
            "C<sub>p</sub> = D⁻/(D⁺ + D⁻). Posisi dengan C<sub>p</sub> tertinggi = posisi paling sesuai. "
            "<b>Lampiran</b> berisi TOPSIS pembanding per posisi (alternatif = pemain) untuk melihat pemain terbaik di tiap posisi.", S["p"]),
          Spacer(1, 4), _legend(criteria, S)]
    pts = [sppk.position_topsis(profiles, p["skills"]) for p in players]
    summ = [["Pemain", "Nama", "Posisi terdaftar"] + [f"Ci {x}" for x in POSITIONS] + ["Rekomendasi", "Skor"]]
    for pl, pt in zip(players, pts):
        cis = [r["ci"] for r in pt["res"]]
        b = cis.index(max(cis))
        summ.append([pl["code"], pl["name"], POSITIONS[pl["position"]]] + [f"{c:.4f}" for c in cis] +
                    [POSITIONS[b], f"{sppk.weighted_scores(profiles, pl['skills'])[b]:.2f}"])
    t = _table(summ, font=7.5)
    t.setStyle(TableStyle([("ALIGN", (1, 0), (2, -1), "LEFT")]))
    el += [P("Ringkasan nilai preferensi dan rekomendasi posisi (TOPSIS per pemain)", S["h2"]), t]

    for pl, pt in zip(players, pts):
        best = pt["order"][0]
        el += [PageBreak(), P(f"TOPSIS per pemain — {pl['code']} · {pl['name']} (rekomendasi: {POSITIONS[best]})", S["h2"]),
               P("Skill pemain dan bobot kriteria w = s/Σs", S["h3"]),
               _table([[""] + cc, ["Skill s"] + pt["skills"], ["w"] + [f"{x:.4f}" for x in pt["w"]]]),
               P("Langkah 1. Matriks keputusan X (bobot AHP tiap posisi)", S["h3"]),
               _table([[""] + cc] + [[POSITIONS[i]] + [f"{v:.4f}" for v in r] for i, r in enumerate(pt["x"])] +
                      [["√Σx²"] + [f"{d:.4f}" for d in pt["den"]]], hl_rows=(len(POSITIONS) + 1,)),
               P("Langkah 2. Matriks ternormalisasi R", S["h3"]),
               _table([[""] + cc] + [[POSITIONS[i]] + [f"{v:.4f}" for v in r] for i, r in enumerate(pt["norm"])]),
               P("Langkah 3. Matriks terbobot V dan solusi ideal", S["h3"]),
               _table([[""] + cc] + [[POSITIONS[i]] + [f"{v:.4f}" for v in r] for i, r in enumerate(pt["v"])] +
                      [["A⁺"] + [f"{v:.4f}" for v in pt["a_pos"]], ["A⁻"] + [f"{v:.4f}" for v in pt["a_neg"]]], hl_rows=(len(POSITIONS) + 1, len(POSITIONS) + 2)),
               P("Langkah 4–5. Jarak ke solusi ideal dan nilai preferensi", S["h3"]),
               _table([["Rank", "Posisi", "D⁺", "D⁻", "Ci"]] +
                      [[pt["rank"][i], POSITIONS[i], f"{pt['res'][i]['d_plus']:.4f}", f"{pt['res'][i]['d_minus']:.4f}", f"{pt['res'][i]['ci']:.4f}"]
                       for i in pt["order"]], font=7, first_col=False, col_widths=[1.5 * cm, 3 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm])]
        if pl is players[0]:
            el += [P(f"Contoh perhitungan: {pl['code']}, posisi {POSITIONS[best]}, kriteria {criteria[0]['name']}", S["h3"]),
                   _example_table(sppk.position_matrix_example(names, pt, best, 0) + sppk.topsis_example(names, POSITIONS, pt, best, 0), S)]

    steps = [sppk.topsis_steps([p["skills"] for p in players], sppk.ahp_pos(profiles, k)["weights"]) for k in range(len(POSITIONS))]
    for k, pos in enumerate(POSITIONS):
        st = steps[k]
        el += [PageBreak(), P(f"Lampiran — Pembanding: ranking pemain di posisi {pos}", S["h2"]),
               P("Bobot AHP", S["h3"]), _table([[""] + cc, ["w"] + [f"{x:.4f}" for x in st["w"]]]),
               P("Langkah 1. Matriks keputusan X", S["h3"]),
               _table([[""] + cc] + [[pc[i]] + r for i, r in enumerate(st["x"])] + [["√Σx²"] + [f"{d:.3f}" for d in st["den"]]],
                      hl_rows=(len(pc) + 1,)),
               P("Langkah 2. Matriks ternormalisasi R", S["h3"]),
               _table([[""] + cc] + [[pc[i]] + [f"{v:.4f}" for v in r] for i, r in enumerate(st["norm"])]),
               PageBreak(),
               P(f"Lampiran — {pos} (lanjutan)", S["h2"]),
               P("Langkah 3. Matriks terbobot V dan solusi ideal", S["h3"]),
               _table([[""] + cc] + [[pc[i]] + [f"{v:.4f}" for v in r] for i, r in enumerate(st["v"])] +
                      [["A⁺"] + [f"{v:.4f}" for v in st["a_pos"]], ["A⁻"] + [f"{v:.4f}" for v in st["a_neg"]]],
                      hl_rows=(len(pc) + 1, len(pc) + 2))]
        rk = [["Rank", "Pemain", "D⁺", "D⁻", "Ci"]] + [[st["rank"][i], pc[i], f"{st['res'][i]['d_plus']:.4f}", f"{st['res'][i]['d_minus']:.4f}",
                                                        f"{st['res'][i]['ci']:.4f}"] for i in st["order"]]
        el += [P("Langkah 4–5. Jarak ke solusi ideal dan nilai preferensi (urut ranking)", S["h3"]),
               _table(rk, font=7, first_col=False, col_widths=[1.5 * cm, 2.5 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm])]
    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


# ===================================================================== EXCEL
HFILL = PatternFill("solid", fgColor="E1EFEA")
AFILL = PatternFill("solid", fgColor="F6EAD9")
THIN = Border(*(Side(style="thin", color="B8C7C2"),) * 4)
BOLD = Font(bold=True)


class Sheet:
    def __init__(self, ws):
        self.ws, self.r = ws, 1

    def title(self, text, size=13):
        self.ws.cell(self.r, 1, text).font = Font(bold=True, size=size, color="0F6B57")
        self.r += 1

    def text(self, text):
        self.ws.cell(self.r, 1, text).alignment = Alignment(wrap_text=False)
        self.r += 1

    def table(self, rows, hl=(), fmt="0.0000"):
        for k, row in enumerate(rows):
            for j, v in enumerate(row):
                c = self.ws.cell(self.r, j + 1, v)
                c.border = THIN
                if isinstance(v, float):
                    c.number_format = fmt
                if k == 0:
                    c.font, c.fill = BOLD, HFILL
                elif j == 0:
                    c.font = BOLD
                if k in hl:
                    c.font, c.fill = BOLD, AFILL
            self.r += 1
        self.r += 1

    def widths(self, first=22, rest=11):
        from openpyxl.utils import get_column_letter
        self.ws.column_dimensions["A"].width = first
        for j in range(2, 30):
            self.ws.column_dimensions[get_column_letter(j)].width = rest


def _wb_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def ahp_xlsx(criteria, profiles):
    wb = Workbook()
    cc = codes(criteria)
    sh = Sheet(wb.active)
    wb.active.title = "Ringkasan"
    sh.title("Perhitungan AHP — Ringkasan Bobot", 14)
    sh.text("Model hibrida AHP–TOPSIS + analisis gap berbobot. Matriks kepentingan ditetapkan sekali oleh pakar per posisi (skala Saaty).")
    sh.text("Matriks berlabel 'Contoh' = nilai awal simulasi, wajib diganti penilaian pakar. Konsisten bila CR < 0,1.")
    sh.r += 1
    al = [sppk.ahp_pos(profiles, p) for p in range(len(POSITIONS))]
    sh.table([["Kode", "Kriteria"] + POSITIONS] + [[cc[i], c["name"]] + [a["weights"][i] for a in al] for i, c in enumerate(criteria)] +
             [["", "λmax"] + [a["lambda_max"] for a in al], ["", "CI"] + [a["ci"] for a in al], ["", "RI"] + [a["ri"] for a in al],
              ["", "CR"] + [a["cr"] for a in al], ["", "Status"] + ["Konsisten" if a["consistent"] else "Tidak" for a in al],
              ["", "Sumber matriks"] + [{"pakar": "Penilaian pakar", "contoh": "Contoh (belum diisi pakar)"}.get(getattr(profiles, "sources", {}).get(k), "Belum diisi (semua = 1)") for k in range(len(al))]])
    sh.widths(8, 16)
    sh.ws.column_dimensions["B"].width = 22
    for p, pos in enumerate(POSITIONS):
        a, s, n = al[p], profiles[p], al[p]["n"]
        sh = Sheet(wb.create_sheet(sheet_name(pos, wb.sheetnames)))
        sh.title(f"AHP — {pos}", 14)
        sh.title("Hasil: bobot prioritas kriteria (vektor prioritas AHP)", 11)
        sh.table([["Kode", "Kriteria", "Bobot", "%", "Peringkat"]] +
                 [[cc[i], c["name"], a["weights"][i], round(a["weights"][i] * 100, 1),
                   1 + sum(1 for x in a["weights"] if x > a["weights"][i])] for i, c in enumerate(criteria)])
        sh.title("Langkah 1. Matriks perbandingan berpasangan — sumber: " + {"pakar": "Penilaian pakar", "contoh": "Contoh (belum diisi pakar)"}.get(getattr(profiles, "sources", {}).get(p), "Belum diisi (semua = 1)"), 11)
        sh.table([[""] + cc] + [[cc[i]] + [round(v, 4) for v in row] for i, row in enumerate(a["matrix"])] +
                 [["Σ kolom"] + a["col_sum"]], hl=(n + 1,))
        sh.title("Langkah 2. Matriks ternormalisasi & bobot", 11)
        sh.table([[""] + cc + ["Bobot w"]] + [[cc[i]] + row + [a["weights"][i]] for i, row in enumerate(a["normalized"])])
        sh.title("Langkah 3. Uji konsistensi", 11)
        sh.table([["Kode", "w", "(Aw)", "(Aw)/w"]] + [[cc[i], a["weights"][i], a["aw"][i], a["ratio"][i]] for i in range(n)])
        sh.table([["Parameter", "Nilai"], ["n", n], ["λmax", a["lambda_max"]], ["CI", a["ci"]], ["RI", a["ri"]], ["CR", a["cr"]],
                  ["Status", "Konsisten" if a["consistent"] else "Tidak konsisten"]])
        sh.title(f"Contoh perhitungan ({criteria[0]['name']} vs {criteria[1]['name'] if n > 1 else criteria[0]['name']})", 11)
        sh.table([["Langkah", "Rumus", "Substitusi angka"]] +
                 [list(x) for x in sppk.ahp_example([c["name"] for c in criteria], s, a, 0, 1 if n > 1 else 0)])
        sh.widths(14, 11)
        sh.ws.column_dimensions["B"].width = 20
    return _wb_bytes(wb)


def topsis_xlsx(criteria, profiles, players):
    wb = Workbook()
    cc = codes(criteria)
    pc = [p["code"] for p in players]
    names = [c["name"] for c in criteria]
    steps = [sppk.topsis_steps([p["skills"] for p in players], sppk.ahp_pos(profiles, k)["weights"]) for k in range(len(POSITIONS))]
    pts = [sppk.position_topsis(profiles, p["skills"]) for p in players]
    sh = Sheet(wb.active)
    wb.active.title = "Ringkasan"
    sh.title("Perhitungan TOPSIS — Ringkasan Nilai Preferensi", 14)
    sh.text("Penentu rekomendasi: TOPSIS per pemain, alternatif = semua posisi, X = bobot AHP tiap posisi, w = skill/Σskill, benefit.")
    sh.text("Ci = D⁻/(D⁺+D⁻). Sheet 'Lampiran ...' = pembanding per posisi (alternatif = pemain aktif).")
    sh.r += 1
    rows = [["Pemain", "Nama", "Posisi terdaftar"] + [f"Ci {x}" for x in POSITIONS] + ["Rekomendasi", "Skor (1-10)"]]
    for pl, pt in zip(players, pts):
        cis = [r["ci"] for r in pt["res"]]
        b = cis.index(max(cis))
        rows.append([pl["code"], pl["name"], POSITIONS[pl["position"]]] + cis + [POSITIONS[b], round(sppk.weighted_scores(profiles, pl["skills"])[b], 2)])
    sh.table(rows)
    sh.title("Keterangan kode kriteria", 11)
    sh.table([["Kode", "Kriteria"]] + [[cc[i], c["name"]] for i, c in enumerate(criteria)])
    sh.widths(10, 16)
    sh.ws.column_dimensions["B"].width = 22
    # satu sheet berisi langkah TOPSIS per pemain
    sh = Sheet(wb.create_sheet("Per Pemain"))
    for pl, pt in zip(players, pts):
        best = pt["order"][0]
        sh.title(f"{pl['code']} · {pl['name']} — rekomendasi: {POSITIONS[best]}", 13)
        sh.table([[""] + cc, ["Skill s"] + pt["skills"], ["w = s/Σs"] + pt["w"]])
        sh.title("Langkah 1. Matriks keputusan X (bobot AHP tiap posisi)", 11)
        sh.table([[""] + cc] + [[POSITIONS[i]] + r for i, r in enumerate(pt["x"])] + [["√Σx²"] + pt["den"]], hl=(len(POSITIONS) + 1,))
        sh.title("Langkah 2. Matriks ternormalisasi R", 11)
        sh.table([[""] + cc] + [[POSITIONS[i]] + r for i, r in enumerate(pt["norm"])])
        sh.title("Langkah 3. Matriks terbobot V & solusi ideal", 11)
        sh.table([[""] + cc] + [[POSITIONS[i]] + r for i, r in enumerate(pt["v"])] + [["A⁺"] + pt["a_pos"], ["A⁻"] + pt["a_neg"]], hl=(len(POSITIONS) + 1, len(POSITIONS) + 2))
        sh.title("Langkah 4–5. Jarak & nilai preferensi", 11)
        sh.table([["Posisi", "D⁺", "D⁻", "Ci", "Rank"]] + [[POSITIONS[i], pt["res"][i]["d_plus"], pt["res"][i]["d_minus"], pt["res"][i]["ci"],
                                                            pt["rank"][i]] for i in range(len(POSITIONS))])
        if pl is players[0]:
            sh.title(f"Contoh perhitungan ({pl['code']}, posisi {POSITIONS[best]}, kriteria {criteria[0]['name']})", 11)
            sh.table([["Langkah", "Rumus", "Substitusi angka"]] +
                     [list(x) for x in sppk.position_matrix_example(names, pt, best, 0) + sppk.topsis_example(names, POSITIONS, pt, best, 0)])
        sh.r += 1
    sh.widths(16, 10)
    for k, pos in enumerate(POSITIONS):
        st = steps[k]
        n = len(pc)
        sh = Sheet(wb.create_sheet(sheet_name("Lampiran " + pos, wb.sheetnames)))
        sh.title(f"Lampiran — Pembanding: ranking pemain di posisi {pos}", 14)
        sh.title("Bobot AHP", 11)
        sh.table([[""] + cc, ["w"] + st["w"]])
        sh.title("Langkah 1. Matriks keputusan X", 11)
        sh.table([[""] + cc] + [[pc[i]] + r for i, r in enumerate(st["x"])] + [["√Σx²"] + st["den"]], hl=(n + 1,), fmt="0.000")
        sh.title("Langkah 2. Matriks ternormalisasi R", 11)
        sh.table([[""] + cc] + [[pc[i]] + r for i, r in enumerate(st["norm"])])
        sh.title("Langkah 3. Matriks terbobot V & solusi ideal", 11)
        sh.table([[""] + cc] + [[pc[i]] + r for i, r in enumerate(st["v"])] + [["A⁺"] + st["a_pos"], ["A⁻"] + st["a_neg"]], hl=(n + 1, n + 2))
        sh.title("Langkah 4–5. Jarak & nilai preferensi", 11)
        sh.table([["Pemain", "D⁺", "D⁻", "Ci", "Rank"]] + [[pc[i], st["res"][i]["d_plus"], st["res"][i]["d_minus"], st["res"][i]["ci"], st["rank"][i]]
                                                       for i in range(n)])
        best = st["order"][0]
        sh.title(f"Contoh perhitungan ({pc[best]}, kriteria {criteria[0]['name']})", 11)
        sh.table([["Langkah", "Rumus", "Substitusi angka"]] +
                 [list(x) for x in sppk.topsis_example([c["name"] for c in criteria], pc, st, best, 0)])
        sh.widths(10, 10)
    return _wb_bytes(wb)
