"""SPK Hibrida AHP–TOPSIS untuk Penentuan Posisi dan Rekomendasi Latihan Pemain Futsal — Flask + SQLite."""
import csv
import io
import json
from datetime import datetime
from functools import wraps

from flask import (Flask, Response, abort, flash, redirect, render_template, request, send_file,
                   session, url_for)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.security import check_password_hash, generate_password_hash

import reports
import sppk
from db import close_db, get_db, init_db
from sppk import POSITIONS

app = Flask(__name__)
app.secret_key = "ganti-dengan-secret-key-acak"
app.teardown_appcontext(close_db)


@app.before_request
def load_positions():
    names = [r["name"] for r in get_db().execute("SELECT name FROM positions ORDER BY idx")]
    POSITIONS[:] = names or POSITIONS


@app.context_processor
def inject():
    u = current_user()
    return {"POSITIONS": POSITIONS, "user": u, "category": sppk.category, "ROLE_LABEL": ROLE_LABEL,
            "is_staff": bool(u and u["role"] in STAFF), "is_super": bool(u and u["role"] == "superadmin"),
            "pending_count": get_db().execute("SELECT COUNT(*) FROM users WHERE status='pending'").fetchone()[0]
            if u and u["role"] == "superadmin" else 0}


@app.template_filter("f")
def fmt(v, d=3):
    return f"{v:.{d}f}"


# ---------------------------------------------------------------- auth & log
STAFF = ("pelatih", "superadmin")
ROLE_LABEL = {"superadmin": "Super Admin", "pelatih": "Pelatih", "pemain": "Pemain"}


def current_user():
    uid = session.get("uid")
    if not uid:
        return None
    return get_db().execute("SELECT * FROM users WHERE id=? AND status='active'", (uid,)).fetchone()


def home_url(u):
    if u["role"] in STAFF:
        return url_for("dashboard")
    return url_for("player_detail", pid=u["player_id"]) if u["player_id"] else url_for("index")


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not current_user():
            return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return w


def coach_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if current_user()["role"] not in STAFF:
            abort(403)
        return f(*a, **k)
    return w


def superadmin_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if current_user()["role"] != "superadmin":
            abort(403)
        return f(*a, **k)
    return w


def log(category, action):
    u = current_user()
    db = get_db()
    db.execute("INSERT INTO logs(username,role,category,action) VALUES(?,?,?,?)",
               (u["username"] if u else "anon", u["role"] if u else "-", category, action))
    db.commit()


def can_view_player(pid):
    u = current_user()
    return u["role"] in STAFF or u["player_id"] == pid


# ---------------------------------------------------------------- data model
def load_model(include_archived=False):
    db = get_db()
    criteria = [dict(r) for r in db.execute("SELECT * FROM criteria ORDER BY sort_order, id")]
    trainings = {}
    for t in db.execute("SELECT * FROM trainings ORDER BY CASE level WHEN 'Dasar' THEN 0 WHEN 'Menengah' THEN 1 ELSE 2 END, sort_order, id"):
        trainings.setdefault(t["criterion_id"], []).append(dict(t))
    for c in criteria:
        c["trainings"] = trainings.get(c["id"], [])
        if c["trainings"]:  # teks ringkas untuk export & snapshot
            c["training"] = "; ".join(t["name"] for t in c["trainings"])
    profiles = sppk.Profiles({p: [] for p in range(len(POSITIONS))})
    prof = {(r["criterion_id"], r["position"]): r["score"] for r in db.execute("SELECT * FROM position_profile")}
    for c in criteria:
        for p in profiles:
            profiles[p].append(prof.get((c["id"], p), 5))
    # matriks kepentingan dari pakar (tersimpan); pasangan yang belum dinilai (mis. kriteria baru) = 1
    pw = {(r["position"], r["ci"], r["cj"]): r["value"] for r in db.execute("SELECT * FROM pairwise")}
    for p in range(len(POSITIONS)):
        n = len(criteria)
        m = [[1.0] * n for _ in range(n)]
        for i, a in enumerate(criteria):
            for j, b in enumerate(criteria):
                if i < j and (p, a["id"], b["id"]) in pw:
                    m[i][j] = pw[(p, a["id"], b["id"])]
                    m[j][i] = 1 / m[i][j]
        profiles.matrices[p] = m
        profiles.sources[p] = ahp_mode_info(p)["source"]
    q = "SELECT * FROM players" + ("" if include_archived else " WHERE archived=0") + " ORDER BY code"
    players = [dict(r) for r in db.execute(q)]
    sk = {(r["player_id"], r["criterion_id"]): r["value"] for r in db.execute("SELECT * FROM skills")}
    for pl in players:
        pl["skills"] = [sk.get((pl["id"], c["id"]), 1) for c in criteria]
        pl["alt_list"] = [int(x) for x in pl["alt_positions"].split(",") if x.strip() != ""]
    return criteria, profiles, players


def get_player(pid):
    pl = get_db().execute("SELECT * FROM players WHERE id=?", (pid,)).fetchone()
    if not pl:
        abort(404)
    return pl


def run_analysis(pid, target):
    criteria, profiles, players = load_model()
    if not any(p["id"] == pid for p in players):
        return None, criteria, None
    if len(players) < 2:
        return None, criteria, None
    res = sppk.analyze(criteria, profiles, players, pid, target)
    pl = next(p for p in players if p["id"] == pid)
    return res, criteria, pl


# ---------------------------------------------------------------- routes: auth
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = get_db().execute("SELECT * FROM users WHERE username=?", (request.form["username"].strip().lower(),)).fetchone()
        if u and check_password_hash(u["password_hash"], request.form["password"]):
            if u["status"] == "pending":
                flash("Akun Anda masih menunggu persetujuan Super Admin.", "warning")
            elif u["status"] == "rejected":
                flash("Pendaftaran akun Anda ditolak. Hubungi Super Admin.", "danger")
            else:
                session["uid"] = u["id"]
                log("akses", "Login")
                return redirect(request.args.get("next") or home_url(u))
        else:
            flash("Username atau password salah.", "danger")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        f = request.form
        uname = f["username"].strip().lower()
        req_role = f.get("requested_role") if f.get("requested_role") in ("pemain", "pelatih") else "pemain"
        if len(f["password"]) < 6 or f["password"] != f["password2"]:
            flash("Password minimal 6 karakter dan kedua isian harus sama.", "danger")
        elif not uname.isalnum():
            flash("Username hanya boleh huruf dan angka.", "danger")
        else:
            db = get_db()
            try:
                db.execute("INSERT INTO users(username,password_hash,role,full_name,status,requested_role,note) VALUES(?,?,?,?,?,?,?)",
                           (uname, generate_password_hash(f["password"]), "pemain", f["full_name"].strip(), "pending",
                            req_role, f.get("note", "").strip()))
                db.execute("INSERT INTO logs(username,role,category,action) VALUES(?,?,?,?)",
                           (uname, "-", "akses", f"Pendaftaran akun baru sebagai {req_role} (menunggu approval)"))
                db.commit()
                flash("Pendaftaran terkirim. Anda bisa login setelah akun disetujui Super Admin.", "success")
                return redirect(url_for("login"))
            except Exception:
                db.rollback()
                flash("Username sudah dipakai.", "danger")
    return render_template("register.html")


@app.route("/logout")
def logout():
    if current_user():
        log("akses", "Logout")
    session.clear()
    return redirect(url_for("index"))


@app.route("/")
def index():
    """Beranda publik: informasi umum tanpa data pribadi pemain."""
    criteria, profiles, players = load_model()
    per_pos = [sum(1 for p in players if p["position"] == i) for i in range(len(POSITIONS))]
    multi = sum(1 for p in players if p["alt_list"])
    ahps = [sppk.ahp_pos(profiles, p) for p in range(len(POSITIONS))]
    top = []
    for a in ahps:
        order = sorted(range(len(criteria)), key=lambda i: a["weights"][i], reverse=True)[:3]
        top.append([(criteria[i]["name"], a["weights"][i]) for i in order])
    return render_template("home.html", criteria=criteria, n_players=len(players), per_pos=per_pos, multi=multi,
                           ahps=ahps, top=top)


@app.route("/dashboard")
@coach_required
def dashboard():
    criteria, profiles, players = load_model()
    all_ci = sppk.topsis_all(profiles, players) if len(players) >= 2 else {}
    rows = []
    for pl in players:
        cis = all_ci.get(pl["id"], [0] * len(POSITIONS))
        best = max(range(len(cis)), key=lambda i: cis[i])
        rows.append({**pl, "cis": cis, "best": best, "score": sppk.weighted_scores(profiles, pl["skills"])[best]})
    db = get_db()
    stats = {"players": len(players), "archived": db.execute("SELECT COUNT(*) FROM players WHERE archived=1").fetchone()[0],
             "criteria": len(criteria), "results": db.execute("SELECT COUNT(*) FROM results").fetchone()[0]}
    ahp_cr = [sppk.ahp_pos(profiles, p)["cr"] for p in range(len(POSITIONS))]
    logs = db.execute("SELECT * FROM logs ORDER BY id DESC LIMIT 8").fetchall()
    return render_template("dashboard.html", rows=rows, stats=stats, ahp_cr=ahp_cr, logs=logs)


# ---------------------------------------------------------------- routes: pemain
@app.route("/players")
@coach_required
def players():
    show = request.args.get("arsip") == "1"
    _, _, pls = load_model(include_archived=True)
    pls = [p for p in pls if bool(p["archived"]) == show]
    return render_template("players.html", players=pls, show_archived=show)


def save_player_form(pid=None):
    f = request.form
    db = get_db()
    criteria, _, _ = load_model()
    alt = ",".join(x for x in f.getlist("alt_positions") if x != f["position"])
    vals = (f["code"].strip().upper(), f["name"].strip(), int(f["position"]), alt,
            int(f["desired_position"]), f.get("assessment", "Tes objektif"), f.get("notes", ""))
    try:
        if pid is None:
            pid = db.execute("INSERT INTO players(code,name,position,alt_positions,desired_position,assessment,notes) "
                             "VALUES(?,?,?,?,?,?,?)", vals).lastrowid
        else:
            db.execute("UPDATE players SET code=?,name=?,position=?,alt_positions=?,desired_position=?,assessment=?,notes=?,"
                       "updated_at=CURRENT_TIMESTAMP WHERE id=?", vals + (pid,))
    except Exception:
        db.rollback()
        flash("Kode pemain sudah dipakai.", "danger")
        return None
    changes = []
    old = {r["criterion_id"]: r["value"] for r in db.execute("SELECT * FROM skills WHERE player_id=?", (pid,))}
    for c in criteria:
        v = max(1, min(10, int(f.get(f"skill_{c['id']}", 5) or 5)))
        if old.get(c["id"]) != v:
            changes.append(f"{c['name']} {old.get(c['id'], '-')}→{v}")
        db.execute("INSERT OR REPLACE INTO skills VALUES(?,?,?)", (pid, c["id"], v))
    db.commit()
    return pid, changes


@app.route("/players/new", methods=["GET", "POST"])
@coach_required
def player_new():
    criteria, _, _ = load_model()
    if request.method == "POST":
        r = save_player_form()
        if r:
            log("input", f"Tambah pemain {request.form['code'].upper()} ({request.form['name']})")
            flash("Pemain ditambahkan.", "success")
            return redirect(url_for("player_detail", pid=r[0]))
    return render_template("player_form.html", player=None, criteria=criteria, skills={})


@app.route("/players/<int:pid>/edit", methods=["GET", "POST"])
@coach_required
def player_edit(pid):
    pl = get_player(pid)
    criteria, _, _ = load_model()
    if request.method == "POST":
        r = save_player_form(pid)
        if r:
            log("input", f"Ubah pemain {request.form['code'].upper()}" + (f": {', '.join(r[1])}" if r[1] else ""))
            flash("Data pemain diperbarui.", "success")
            return redirect(url_for("player_detail", pid=pid))
    skills = {r["criterion_id"]: r["value"] for r in get_db().execute("SELECT * FROM skills WHERE player_id=?", (pid,))}
    return render_template("player_form.html", player=pl, criteria=criteria, skills=skills)


@app.route("/players/<int:pid>/archive", methods=["POST"])
@coach_required
def player_archive(pid):
    pl = get_player(pid)
    db = get_db()
    to = 0 if pl["archived"] else 1
    db.execute("UPDATE players SET archived=?, archived_at=? WHERE id=?", (to, datetime.now().isoformat(" ", "seconds") if to else None, pid))
    db.commit()
    log("input", f"{'Arsipkan' if to else 'Pulihkan'} pemain {pl['code']}")
    flash(f"Pemain {pl['code']} {'dipindahkan ke arsip' if to else 'dipulihkan'}.", "success")
    return redirect(url_for("players", arsip="0" if to else "1"))


@app.route("/players/<int:pid>")
@login_required
def player_detail(pid):
    if not can_view_player(pid):
        abort(403)
    pl = get_player(pid)
    target = int(request.args.get("target", pl["desired_position"] if pl["desired_position"] is not None else pl["position"]))
    res, criteria, plm = run_analysis(pid, target)
    if res:
        log("perhitungan", f"Hitung AHP-TOPSIS {pl['code']}: terbaik {res['best']['name']} (Ci={res['best']['ci']:.3f}), target {res['target_name']}")
    history = get_db().execute("SELECT * FROM results WHERE player_id=? ORDER BY id DESC LIMIT 10", (pid,)).fetchall()
    pt = sppk.position_topsis(load_model()[1], plm["skills"]) if res else None
    return render_template("player_detail.html", p=pl, pm=plm, res=res, criteria=criteria, target=target, history=history, pt=pt)


@app.route("/players/<int:pid>/save", methods=["POST"])
@coach_required
def save_result(pid):
    pl = get_player(pid)
    target = int(request.form["target"])
    res, _, plm = run_analysis(pid, target)
    if not res:
        abort(400)
    payload = {"skills": plm["skills"], **{k: v for k, v in res.items() if k != "pt"}}
    db = get_db()
    db.execute("INSERT INTO results(player_id,target_position,best_position,best_ci,payload,created_by) VALUES(?,?,?,?,?,?)",
               (pid, target, res["best"]["pos"], res["best"]["ci"], json.dumps(payload), current_user()["username"]))
    db.commit()
    log("rekomendasi", f"Simpan hasil & rekomendasi {pl['code']}: {len(res['recommendations'])} training untuk target {res['target_name']}")
    flash("Hasil perhitungan dan rekomendasi disimpan.", "success")
    return redirect(url_for("player_detail", pid=pid, target=target))


# ---------------------------------------------------------------- export
def analysis_rows(pl, res):
    rows = [["Pemain", f"{pl['code']} - {pl['name']}"], ["Posisi terdaftar", POSITIONS[pl["position"]]],
            ["Posisi paling sesuai", res["best"]["name"]], ["Posisi diinginkan", res["target_name"]],
            ["Kesiapan posisi target", f"{res['target_readiness'] * 100:.1f}%"], [],
            ["RANKING POSISI (TOPSIS) + SKOR PENDUKUNG Σ(bobot × nilai)"], ["Rank", "Posisi", "D+", "D-", "Ci", "Skor (1-10)"]]
    rows += [[i + 1, r["name"], round(r["d_plus"], 4), round(r["d_minus"], 4), round(r["ci"], 4), round(r["score"], 2)] for i, r in enumerate(res["ranking"])]
    rows += [[], ["GAP SKILL & REKOMENDASI TRAINING"], ["Prioritas", "Kriteria", "Skill", "Target", "Gap", "Bobot AHP", "Skor prioritas", "Training"]]
    rows += [[i + 1, g["criterion"], g["skill"], g["target"], g["gap"], round(g["weight"], 4), round(g["priority"], 4),
              g["training"] if g["gap"] else "-"] for i, g in enumerate(res["gaps"])]
    return rows


def as_csv(rows, name):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return Response("﻿" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={name}.csv"})


def as_xlsx(sheets, name):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            if row and isinstance(row[0].value, str) and (row[0].value.isupper() or row[0].value in ("Rank", "Prioritas", "Kode", "Waktu")):
                for c in row:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor="E1EFEA")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{name}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/players/<int:pid>/export.<fmt>")
@login_required
def export_player(pid, fmt):
    if not can_view_player(pid):
        abort(403)
    pl = get_player(pid)
    target = int(request.args.get("target", pl["desired_position"] or 0))
    res, _, _ = run_analysis(pid, target)
    if not res:
        abort(400)
    rows = analysis_rows(pl, res)
    log("rekomendasi", f"Export {fmt.upper()} hasil {pl['code']}")
    name = f"hasil_{pl['code']}_{datetime.now():%Y%m%d}"
    return as_csv(rows, name) if fmt == "csv" else as_xlsx({"Hasil": rows}, name)


@app.route("/topsis/export.<fmt>")
@coach_required
def export_topsis(fmt):
    criteria, profiles, players = load_model()
    all_ci = sppk.topsis_all(profiles, players)
    rows = [["Kode", "Nama", "Posisi terdaftar"] + [f"Ci {p}" for p in POSITIONS] + ["Rekomendasi"]]
    for pl in players:
        cis = all_ci[pl["id"]]
        rows.append([pl["code"], pl["name"], POSITIONS[pl["position"]]] + [round(c, 4) for c in cis] + [POSITIONS[cis.index(max(cis))]])
    sheets = {"Ranking TOPSIS": rows}
    for p, pos in enumerate(POSITIONS):
        a = sppk.ahp_pos(profiles, p)
        sheets[reports.sheet_name(f"AHP {pos}", sheets)] = [["Kode", "Kriteria", "Bobot", "%", "Peringkat"]] + \
            [[f"C{i + 1}", c["name"], round(a["weights"][i], 4), round(a["weights"][i] * 100, 1),
              1 + sum(1 for x in a["weights"] if x > a["weights"][i])] for i, c in enumerate(criteria)] + \
            [[], ["Lambda max", round(a["lambda_max"], 4)], ["CI", round(a["ci"], 4)], ["RI", a["ri"]], ["CR", round(a["cr"], 4)]]
    sheets["Skill Pemain"] = [["Kode", "Nama"] + [c["name"] for c in criteria]] + [[p["code"], p["name"]] + p["skills"] for p in players]
    log("perhitungan", f"Export {fmt.upper()} ranking TOPSIS semua pemain")
    name = f"ranking_topsis_{datetime.now():%Y%m%d}"
    return as_csv(rows, name) if fmt == "csv" else as_xlsx(sheets, name)


# ---------------------------------------------------------------- kriteria & AHP
@app.route("/criteria", methods=["GET", "POST"])
@coach_required
def criteria_view():
    db = get_db()
    if request.method == "POST":
        act = request.form["action"]
        try:
            if act == "add":
                name = request.form["name"].strip()
                tr = request.form.get("training", "").strip() or f"Latihan khusus {name}"
                cid = db.execute("INSERT INTO criteria(name,training,sort_order) VALUES(?,?,(SELECT COALESCE(MAX(sort_order),0)+1 FROM criteria))",
                                 (name, tr)).lastrowid
                db.execute("INSERT INTO trainings(criterion_id,name,duration,level) VALUES(?,?,?,?)", (cid, tr, "20 menit", "Menengah"))
                for p in range(len(POSITIONS)):
                    db.execute("INSERT INTO position_profile VALUES(?,?,?)", (cid, p, int(request.form.get(f"pos_{p}", 5))))
                for (pid,) in db.execute("SELECT id FROM players").fetchall():
                    db.execute("INSERT INTO skills VALUES(?,?,5)", (pid, cid))
                log("input", f"Tambah kriteria {name}")
                flash(f"Kriteria {name} ditambahkan. Skill semua pemain untuk kriteria ini diisi 5 — perbarui di data pemain.", "success")
            elif act == "edit":
                cid = int(request.form["id"])
                db.execute("UPDATE criteria SET name=? WHERE id=?", (request.form["name"].strip(), cid))
                for p in range(len(POSITIONS)):
                    db.execute("INSERT OR REPLACE INTO position_profile VALUES(?,?,?)",
                               (cid, p, max(1, min(9, int(request.form.get(f"pos_{p}", 5) or 5)))))
                log("input", f"Ubah kriteria {request.form['name']}")
                flash("Kriteria diperbarui.", "success")
            elif act == "delete":
                cid = int(request.form["id"])
                if db.execute("SELECT COUNT(*) FROM criteria").fetchone()[0] <= 3:
                    flash("Minimal 3 kriteria diperlukan untuk AHP.", "danger")
                else:
                    name = db.execute("SELECT name FROM criteria WHERE id=?", (cid,)).fetchone()["name"]
                    db.execute("DELETE FROM criteria WHERE id=?", (cid,))
                    log("input", f"Hapus kriteria {name}")
                    flash(f"Kriteria {name} dihapus.", "success")
            db.commit()
        except Exception:
            db.rollback()
            flash("Nama kriteria sudah ada.", "danger")
        return redirect(url_for("criteria_view"))
    criteria, profiles, _ = load_model()
    return render_template("criteria.html", criteria=criteria, profiles=profiles)


@app.route("/positions", methods=["GET", "POST"])
@coach_required
def positions_view():
    db = get_db()
    if request.method == "POST":
        f, act = request.form, request.form["action"]
        try:
            if act == "add":
                name = f["name"].strip()
                n = db.execute("SELECT COUNT(*) FROM positions").fetchone()[0]
                db.execute("INSERT INTO positions(idx,name,description) VALUES(?,?,?)", (n, name, f.get("description", "").strip()))
                for (cid,) in db.execute("SELECT id FROM criteria").fetchall():
                    db.execute("INSERT OR REPLACE INTO position_profile VALUES(?,?,5)", (cid, n))
                log("input", f"Tambah posisi {name}")
                flash(f"Posisi {name} ditambahkan. Matriks kepentingannya masih kosong (semua = 1, bobot sama rata): "
                      "minta pakar mengisinya di Analisis Data → Perhitungan AHP, lalu atur target skor kebutuhan.", "warning")
            elif act == "edit":
                idx = int(f["idx"])
                db.execute("UPDATE positions SET name=?, description=? WHERE idx=?", (f["name"].strip(), f.get("description", "").strip(), idx))
                log("input", f"Ubah posisi #{idx + 1} menjadi {f['name'].strip()}")
                flash("Posisi diperbarui.", "success")
            elif act == "delete":
                idx = int(f["idx"])
                name = db.execute("SELECT name FROM positions WHERE idx=?", (idx,)).fetchone()["name"]
                n = db.execute("SELECT COUNT(*) FROM positions").fetchone()[0]
                used = db.execute("SELECT COUNT(*) FROM players WHERE position=? OR desired_position=?", (idx, idx)).fetchone()[0]
                used_res = db.execute("SELECT COUNT(*) FROM results WHERE target_position=? OR best_position=?", (idx, idx)).fetchone()[0]
                if n <= 2:
                    flash("Minimal 2 posisi diperlukan untuk TOPSIS.", "danger")
                elif used or used_res:
                    flash(f"Posisi {name} tidak bisa dihapus: dipakai {used} pemain (posisi utama/diinginkan) dan {used_res} hasil tersimpan. "
                          "Pindahkan pemain ke posisi lain terlebih dahulu.", "danger")
                else:
                    delete_position(db, idx)
                    log("input", f"Hapus posisi {name}")
                    flash(f"Posisi {name} dihapus.", "success")
            db.commit()
        except Exception:
            db.rollback()
            flash("Nama posisi sudah dipakai.", "danger")
        return redirect(url_for("positions_view"))
    criteria, profiles, players = load_model(include_archived=True)
    rows = []
    for i, name in enumerate(POSITIONS):
        a = sppk.ahp_pos(profiles, i)
        d = db.execute("SELECT description FROM positions WHERE idx=?", (i,)).fetchone()
        rows.append({"idx": i, "name": name, "description": d["description"] if d else "",
                     "main": sum(1 for p in players if p["position"] == i and not p["archived"]),
                     "alt": sum(1 for p in players if i in p["alt_list"] and not p["archived"]),
                     "used": sum(1 for p in players if p["position"] == i or p["desired_position"] == i) +
                     db.execute("SELECT COUNT(*) FROM results WHERE target_position=? OR best_position=?", (i, i)).fetchone()[0],
                     "source": ahp_mode_info(i)["source"], "cr": a["cr"], "consistent": a["consistent"],
                     "top": sorted(zip([c["name"] for c in criteria], a["weights"]), key=lambda x: -x[1])[:3]})
    return render_template("positions.html", rows=rows)


def delete_position(db, idx):
    """Hapus posisi idx lalu geser indeks posisi di atasnya turun satu agar tetap 0..n-1."""
    db.execute("DELETE FROM positions WHERE idx=?", (idx,))
    for t in ("position_profile", "ahp_mode", "pairwise"):
        db.execute(f"DELETE FROM {t} WHERE position=?", (idx,))
    n = db.execute("SELECT COUNT(*) FROM positions").fetchone()[0]
    for k in range(idx + 1, n + 1):  # urut naik: slot di bawahnya sudah kosong
        db.execute("UPDATE positions SET idx=? WHERE idx=?", (k - 1, k))
        for t in ("position_profile", "ahp_mode", "pairwise"):
            db.execute(f"UPDATE {t} SET position=? WHERE position=?", (k - 1, k))
        db.execute("UPDATE players SET position=? WHERE position=?", (k - 1, k))
        db.execute("UPDATE players SET desired_position=? WHERE desired_position=?", (k - 1, k))
        db.execute("UPDATE results SET target_position=? WHERE target_position=?", (k - 1, k))
        db.execute("UPDATE results SET best_position=? WHERE best_position=?", (k - 1, k))
    for pid, alt in db.execute("SELECT id, alt_positions FROM players").fetchall():
        vals = [int(x) for x in (alt or "").split(",") if x.strip() != ""]
        new = [str(v - 1 if v > idx else v) for v in vals if v != idx]
        db.execute("UPDATE players SET alt_positions=? WHERE id=?", (",".join(new), pid))


LEVELS = ["Dasar", "Menengah", "Lanjut"]


@app.route("/trainings", methods=["GET", "POST"])
@coach_required
def trainings_view():
    db = get_db()
    if request.method == "POST":
        f, act = request.form, request.form["action"]
        level = f.get("level") if f.get("level") in LEVELS else "Menengah"
        if act == "add":
            cid = int(f["criterion_id"])
            db.execute("INSERT INTO trainings(criterion_id,name,description,duration,level,sort_order) "
                       "VALUES(?,?,?,?,?,(SELECT COALESCE(MAX(sort_order),0)+1 FROM trainings WHERE criterion_id=?))",
                       (cid, f["name"].strip(), f.get("description", "").strip(), f.get("duration", "").strip(), level, cid))
            log("input", f"Tambah latihan {f['name'].strip()}")
            flash("Latihan ditambahkan.", "success")
        elif act == "edit":
            db.execute("UPDATE trainings SET criterion_id=?, name=?, description=?, duration=?, level=? WHERE id=?",
                       (int(f["criterion_id"]), f["name"].strip(), f.get("description", "").strip(), f.get("duration", "").strip(), level, int(f["id"])))
            log("input", f"Ubah latihan {f['name'].strip()}")
            flash("Latihan diperbarui.", "success")
        elif act == "delete":
            t = db.execute("SELECT * FROM trainings WHERE id=?", (int(f["id"]),)).fetchone()
            left = db.execute("SELECT COUNT(*) FROM trainings WHERE criterion_id=?", (t["criterion_id"],)).fetchone()[0]
            if left <= 1:
                flash("Setiap kriteria minimal memiliki 1 latihan agar rekomendasi tetap tersedia.", "danger")
            else:
                db.execute("DELETE FROM trainings WHERE id=?", (t["id"],))
                log("input", f"Hapus latihan {t['name']}")
                flash(f"Latihan {t['name']} dihapus.", "success")
        db.commit()
        return redirect(url_for("trainings_view", c=f.get("criterion_id", "")))
    criteria, _, _ = load_model()
    sel = request.args.get("c", "")
    return render_template("trainings.html", criteria=criteria, levels=LEVELS, sel=sel)


@app.route("/ahp")
@login_required
def ahp_view():
    pos = int(request.args.get("pos", 0))
    criteria, profiles, _ = load_model()
    a = sppk.ahp_pos(profiles, pos)
    n = len(criteria)
    ei, ej = min(int(request.args.get("ci", 0)), n - 1), min(int(request.args.get("cj", 1)), n - 1)
    example = sppk.ahp_example([c["name"] for c in criteria], profiles[pos], a, ei, ej)
    return render_template("ahp.html", pos=pos, criteria=criteria, scores=profiles[pos], a=a, example=example, ei=ei, ej=ej,
                           summary=[sppk.ahp_pos(profiles, p) for p in range(len(POSITIONS))], mode_info=ahp_mode_info(pos))


def ahp_mode_info(pos):
    """Sumber matriks posisi: 'pakar' (diisi pelatih), 'contoh' (nilai awal, belum diisi pakar), 'kosong'."""
    r = get_db().execute("SELECT * FROM ahp_mode WHERE position=?", (pos,)).fetchone()
    if not r or r["mode"] != "manual":
        return {"source": "kosong", "updated_by": None, "updated_at": None}
    return {"source": "contoh" if r["updated_by"] == "contoh" else "pakar", **dict(r)}


@app.route("/ahp/<int:pos>/matriks", methods=["GET", "POST"])
@coach_required
def ahp_matrix(pos):
    if not 0 <= pos < len(POSITIONS):
        abort(404)
    criteria, profiles, _ = load_model()
    db = get_db()
    if request.method == "POST":
        user = current_user()["username"]
        if request.form.get("action") != "save":
            abort(400)
        n = len(criteria)
        m = [[1.0] * n for _ in range(n)]
        for i in range(n):
            for j in range(i + 1, n):
                v = sppk.parse_saaty(request.form.get(f"a_{i}_{j}", "1"))
                m[i][j], m[j][i] = v, 1 / v
        a = sppk.ahp(profiles[pos], m)
        if not a["consistent"]:
            # ditolak sampai konsisten: tidak ada yang tersimpan, bobot lama tetap dipakai
            log("input", f"Matriks pakar {POSITIONS[pos]} ditolak: CR={a['cr']:.4f} ≥ 0,1")
            flash(f"Belum disimpan: CR = {a['cr']:.4f} ≥ 0,1 (penilaian belum konsisten). "
                  "Perbaiki pasangan pada petunjuk di bawah, lalu cek ulang.", "danger")
            hints = sppk.inconsistency_hints([c["name"] for c in criteria], m, a["weights"])
            return render_template("ahp_matrix.html", pos=pos, criteria=criteria, a=a, mode_info=ahp_mode_info(pos),
                                   saaty=sppk.SAATY, RI=sppk.RI, hints=hints)
        db.execute("DELETE FROM pairwise WHERE position=?", (pos,))
        for i in range(n):
            for j in range(i + 1, n):
                db.execute("INSERT INTO pairwise VALUES(?,?,?,?)", (pos, criteria[i]["id"], criteria[j]["id"], m[i][j]))
        db.execute("INSERT OR REPLACE INTO ahp_mode VALUES(?, 'manual', ?, CURRENT_TIMESTAMP)", (pos, user))
        db.commit()
        log("input", f"Simpan matriks kepentingan pakar {POSITIONS[pos]} (CR={a['cr']:.4f}); bobot & konsistensi dihitung ulang")
        flash(f"Matriks {POSITIONS[pos]} disimpan. CR = {a['cr']:.4f} (konsisten); bobot baru dipakai untuk semua perhitungan.", "success")
        return redirect(url_for("ahp_view", pos=pos))
    a = sppk.ahp_pos(profiles, pos)
    return render_template("ahp_matrix.html", pos=pos, criteria=criteria, a=a, mode_info=ahp_mode_info(pos),
                           saaty=sppk.SAATY, RI=sppk.RI)


@app.route("/topsis")
@coach_required
def topsis_view():
    criteria, profiles, players = load_model()
    names = [c["name"] for c in criteria]
    view = request.args.get("view", "pemain")
    ec = min(int(request.args.get("c", 0)), len(criteria) - 1)
    if view == "posisi":
        # pembanding: alternatif = pemain aktif, per posisi
        pos = int(request.args.get("pos", 0))
        st = sppk.topsis_steps([p["skills"] for p in players], sppk.ahp_pos(profiles, pos)["weights"])
        ep = min(int(request.args.get("p", 0)), len(players) - 1)
        example = sppk.topsis_example(names, [p["code"] for p in players], st, ep, ec)
        return render_template("topsis.html", view=view, pos=pos, criteria=criteria, players=players, st=st, example=example,
                               ep=ep, ec=ec, w=st["w"], norm=st["norm"], v=st["v"], a_pos=st["a_pos"], a_neg=st["a_neg"],
                               res=st["res"], order=st["order"])
    # utama: TOPSIS per pemain, alternatif = posisi
    ep = min(int(request.args.get("p", 0)), len(players) - 1)
    st = sppk.position_topsis(profiles, players[ep]["skills"])
    ex_pos = min(int(request.args.get("ap", st["order"][0])), len(POSITIONS) - 1)
    example = sppk.position_matrix_example(names, st, ex_pos, ec) + sppk.topsis_example(names, POSITIONS, st, ex_pos, ec)
    summary = sppk.topsis_all(profiles, players)
    scores = {pl["id"]: sppk.weighted_scores(profiles, pl["skills"]) for pl in players}
    return render_template("topsis.html", view=view, criteria=criteria, players=players, st=st, example=example,
                           ep=ep, ec=ec, ex_pos=ex_pos, summary=summary, scores=scores)


@app.route("/ahp/export.<fmt>")
@login_required
def export_ahp(fmt):
    criteria, profiles, _ = load_model()
    log("perhitungan", f"Export {fmt.upper()} laporan perhitungan AHP")
    name = f"perhitungan_AHP_{datetime.now():%Y%m%d}"
    if fmt == "pdf":
        return send_file(reports.ahp_pdf(criteria, profiles), download_name=f"{name}.pdf", mimetype="application/pdf")
    return send_file(reports.ahp_xlsx(criteria, profiles), as_attachment=True, download_name=f"{name}.xlsx")


@app.route("/topsis/report.<fmt>")
@coach_required
def export_topsis_report(fmt):
    criteria, profiles, players = load_model()
    if len(players) < 2:
        abort(400)
    log("perhitungan", f"Export {fmt.upper()} laporan perhitungan TOPSIS")
    name = f"perhitungan_TOPSIS_{datetime.now():%Y%m%d}"
    if fmt == "pdf":
        return send_file(reports.topsis_pdf(criteria, profiles, players), download_name=f"{name}.pdf", mimetype="application/pdf")
    return send_file(reports.topsis_xlsx(criteria, profiles, players), as_attachment=True, download_name=f"{name}.xlsx")


# ---------------------------------------------------------------- hasil & log & akun
@app.route("/results")
@login_required
def results():
    u = current_user()
    q = "SELECT r.*, p.code, p.name FROM results r JOIN players p ON p.id=r.player_id"
    rows = get_db().execute(q + (" WHERE r.player_id=?" if u["role"] == "pemain" else "") + " ORDER BY r.id DESC",
                            (u["player_id"],) if u["role"] == "pemain" else ()).fetchall()
    return render_template("results.html", rows=rows)


@app.route("/results/<int:rid>")
@login_required
def result_detail(rid):
    r = get_db().execute("SELECT r.*, p.code, p.name FROM results r JOIN players p ON p.id=r.player_id WHERE r.id=?", (rid,)).fetchone()
    if not r or not can_view_player(r["player_id"]):
        abort(404)
    return render_template("result_detail.html", r=r, res=json.loads(r["payload"]))


@app.route("/logs")
@coach_required
def logs():
    cat = request.args.get("cat", "")
    q = "SELECT * FROM logs" + (" WHERE category=?" if cat else "") + " ORDER BY id DESC LIMIT 500"
    rows = get_db().execute(q, (cat,) if cat else ()).fetchall()
    if request.args.get("export") == "csv":
        return as_csv([["Waktu", "User", "Role", "Kategori", "Aktivitas"]] +
                      [[r["created_at"], r["username"], r["role"], r["category"], r["action"]] for r in rows], "log_aktivitas")
    return render_template("logs.html", rows=rows, cat=cat)


@app.route("/users", methods=["GET", "POST"])
@superadmin_required
def users():
    db = get_db()
    if request.method == "POST":
        f = request.form
        try:
            uid = int(f["id"]) if f.get("id") else None
            uname = db.execute("SELECT username FROM users WHERE id=?", (uid,)).fetchone()["username"] if uid else f.get("username", "")
            if f["action"] in ("add", "approve"):
                role = f["role"] if f["role"] in ROLE_LABEL else "pemain"
                pid = int(f["player_id"]) if f.get("player_id") and role == "pemain" else None
                if role == "pemain" and not pid:
                    flash("Pilih pemain yang ditautkan untuk akun pemain.", "danger")
                    return redirect(url_for("users"))
            if f["action"] == "add":
                if len(f["password"]) < 6:
                    flash("Password minimal 6 karakter.", "danger")
                    return redirect(url_for("users"))
                db.execute("INSERT INTO users(username,password_hash,role,player_id,full_name,status,approved_by,approved_at) "
                           "VALUES(?,?,?,?,?,'active',?,CURRENT_TIMESTAMP)",
                           (uname.strip().lower(), generate_password_hash(f["password"]), role, pid, f.get("full_name", ""),
                            current_user()["username"]))
                log("akses", f"Buat akun {uname} ({role})")
                flash(f"Akun {uname} dibuat.", "success")
            elif f["action"] == "approve":
                db.execute("UPDATE users SET status='active', role=?, player_id=?, approved_by=?, approved_at=CURRENT_TIMESTAMP WHERE id=?",
                           (role, pid, current_user()["username"], uid))
                log("akses", f"Setujui akun {uname} sebagai {role}")
                flash(f"Akun {uname} disetujui sebagai {ROLE_LABEL[role]}.", "success")
            elif f["action"] == "reject":
                db.execute("UPDATE users SET status='rejected', approved_by=?, approved_at=CURRENT_TIMESTAMP WHERE id=?",
                           (current_user()["username"], uid))
                log("akses", f"Tolak pendaftaran akun {uname}")
                flash(f"Pendaftaran {uname} ditolak.", "warning")
            elif f["action"] == "password":
                if len(f["password"]) < 6:
                    flash("Password minimal 6 karakter.", "danger")
                    return redirect(url_for("users"))
                db.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(f["password"]), uid))
                log("akses", f"Ubah password akun {uname}")
                flash(f"Password akun {uname} sekarang: {f['password']} (sampaikan ke pemilik akun; tidak akan ditampilkan lagi).", "success")
            db.commit()
        except Exception:
            db.rollback()
            flash("Username sudah dipakai.", "danger")
        return redirect(url_for("users"))
    rows = db.execute("SELECT u.*, p.code FROM users u LEFT JOIN players p ON p.id=u.player_id WHERE u.status!='pending' "
                      "ORDER BY CASE u.role WHEN 'superadmin' THEN 0 WHEN 'pelatih' THEN 1 ELSE 2 END, u.username").fetchall()
    pending = db.execute("SELECT * FROM users WHERE status='pending' ORDER BY id").fetchall()
    pls = db.execute("SELECT id, code, name FROM players WHERE archived=0 ORDER BY code").fetchall()
    return render_template("users.html", rows=rows, pending=pending, players=pls)


@app.route("/panduan")
def panduan():
    return render_template("panduan.html")


HELP_ANCHOR = {"positions_view": "positions_view", "trainings_view": "trainings_view", "ahp_matrix": "ahp_matrix", "index": "home", "dashboard": "index", "player_new": "player_form", "player_edit": "player_form", "result_detail": "results", "panduan": "alur"}


@app.context_processor
def help_link():
    ep = request.endpoint or ""
    return {"help_anchor": HELP_ANCHOR.get(ep, ep)}


@app.errorhandler(403)
def forbidden(_):
    return render_template("error.html", msg="Anda tidak memiliki akses ke halaman ini."), 403


@app.errorhandler(404)
def not_found(_):
    return render_template("error.html", msg="Halaman atau data tidak ditemukan."), 404


if __name__ == "__main__":
    import sys
    init_db(reset="--reset" in sys.argv)
    app.run(debug=True, port=5000)
